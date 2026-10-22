from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Colors
PRIMARY_FILL = PatternFill(start_color="0A2540", end_color="0A2540", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NORMAL_FONT = Font(color="000000")

def _auto_width(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

def _format_header(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = PRIMARY_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def create_weekly_excel_report(
    department,
    week_start,
    week_end,
    summary,
    insights,
    output_dir="/opt/airflow/reports",
):
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    department_code = department["department_code"]
    department_name = department["department_name"]
    filename = f"{department_code}_Weekly_Report_{week_start}_to_{week_end}.xlsx"
    filepath = output_path / filename

    workbook = Workbook()
    workbook.remove(workbook.active)
    
    # Calculate Overall Status
    dept = summary["department"]
    att_pct = dept["average_attendance_percentage"]
    acad_pct = dept["average_academic_percentage"]
    
    overall_status = "Healthy"
    priority_issues = []
    
    if dept["students_with_attendance_risk"] > 0:
        priority_issues.append(f"{dept['students_with_attendance_risk']} students with high attendance risk (<75%).")
    risk_batches = [b for b in summary.get("batches", []) if b["attendance_concern"] or b["academic_concern"]]
    if risk_batches:
        priority_issues.append(f"{len(risk_batches)} batches flagged for academic or attendance concerns.")
    risk_courses = [c for c in summary.get("courses", []) if c["average_academic_percentage"] is not None and c["average_academic_percentage"] < 50]
    if risk_courses:
        priority_issues.append(f"{len(risk_courses)} courses showing low academic performance (<50%).")
    if len(priority_issues) >= 3 or (att_pct is not None and att_pct < 75):
        overall_status = "Attention Required"
    elif len(priority_issues) > 0:
        overall_status = "Monitor"

    # -----------------------------------------
    # 01 Executive Summary
    # -----------------------------------------
    ws1 = workbook.create_sheet("01 Executive Summary")
    ws1.append(["Metric", "Value"])
    ws1.append(["Department", f"{department_name} ({department_code})"])
    ws1.append(["Reporting Period", f"{week_start} through {week_end}"])
    ws1.append(["Generation Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws1.append(["Overall Status", overall_status])
    ws1.append([])
    ws1.append(["Total Students", dept["student_count"]])
    ws1.append(["Total Courses", dept["course_count"]])
    ws1.append(["Average Attendance %", f"{att_pct}%" if att_pct is not None else "N/A"])
    ws1.append(["Average Academic %", f"{acad_pct}%" if acad_pct is not None else "N/A"])
    ws1.append(["Average Feedback Rating", f"{dept['average_feedback_rating']} / 5" if dept["average_feedback_rating"] is not None else "N/A"])
    ws1.append(["Attendance Risk Students", dept["students_with_attendance_risk"]])
    ws1.append([])
    ws1.append(["Priority Issues"])
    for i, issue in enumerate(priority_issues):
        ws1.append([f"{i+1}. {issue}"])
    
    for row in ws1.iter_rows():
        for cell in row:
            if cell.row == 1 or cell.value == "Priority Issues":
                cell.font = HEADER_FONT
                cell.fill = PRIMARY_FILL
    _auto_width(ws1)

    # -----------------------------------------
    # 02 Course Performance
    # -----------------------------------------
    ws2 = workbook.create_sheet("02 Course Performance")
    ws2.append(["Course Code", "Course Name", "Students", "Total Classes", "Attended Classes", "Attendance %", "Academic %", "Feedback Responses", "Feedback Rating", "Attendance Risk Students"])
    for course in summary["courses"]:
        ws2.append([
            course["course_code"],
            course["course_name"],
            course["student_count"],
            course["total_classes"],
            course["attended_classes"],
            f"{course['average_attendance_percentage']}%" if course['average_attendance_percentage'] is not None else "N/A",
            f"{course['average_academic_percentage']}%" if course['average_academic_percentage'] is not None else "N/A",
            course["feedback_response_count"],
            f"{course['average_feedback_rating']} / 5" if course['average_feedback_rating'] is not None else "N/A",
            course["attendance_risk_students"],
        ])
    _format_header(ws2)
    ws2.freeze_panes = "A2"
    _auto_width(ws2)

    # -----------------------------------------
    # 03 Batch & Section
    # -----------------------------------------
    ws3 = workbook.create_sheet("03 Batch & Section")
    ws3.append(["Type", "Name", "Total Classes", "Attended Classes", "Attendance %", "Academic %", "Attendance Concern (<75%)", "Academic Concern (<50%)"])
    for batch in summary.get("batches", []):
        ws3.append([
            "Batch",
            batch["batch_name"],
            batch["total_classes"],
            batch["attended_classes"],
            f"{batch['average_attendance_percentage']}%" if batch['average_attendance_percentage'] is not None else "N/A",
            f"{batch['average_academic_percentage']}%" if batch['average_academic_percentage'] is not None else "N/A",
            "Yes" if batch["attendance_concern"] else "No",
            "Yes" if batch["academic_concern"] else "No"
        ])
    for sec in summary.get("sections", []):
        ws3.append([
            "Section",
            sec["section_name"],
            sec["total_classes"],
            sec["attended_classes"],
            f"{sec['average_attendance_percentage']}%" if sec['average_attendance_percentage'] is not None else "N/A",
            f"{sec['average_academic_percentage']}%" if sec['average_academic_percentage'] is not None else "N/A",
            "Yes" if sec["attendance_concern"] else "No",
            "Yes" if sec["academic_concern"] else "No"
        ])
    _format_header(ws3)
    ws3.freeze_panes = "A2"
    _auto_width(ws3)

    # -----------------------------------------
    # 04 Attendance Risk
    # -----------------------------------------
    ws4 = workbook.create_sheet("04 Attendance Risk")
    ws4.append(["Student ID", "USN", "Student Name", "Course Code", "Course Name", "Batch", "Section", "Attendance %"])
    raw_data = summary.get("raw_data", [])
    att_risks = [r for r in raw_data if r["data_type"] == "Attendance" and r["metric_3_value"] is not None and float(r["metric_3_value"]) < 75]
    for r in att_risks:
        ws4.append([
            r["student_id"],
            r["usn"],
            r["student_name"],
            r["course_code"],
            r["course_name"],
            r["batch_name"],
            r["section_name"],
            f"{r['metric_3_value']}%"
        ])
    _format_header(ws4)
    ws4.freeze_panes = "A2"
    _auto_width(ws4)

    # -----------------------------------------
    # 05 Academic Risk
    # -----------------------------------------
    ws5 = workbook.create_sheet("05 Academic Risk")
    ws5.append(["Student ID", "USN", "Student Name", "Course Code", "Course Name", "Batch", "Section", "Academic %"])
    acad_risks = [r for r in raw_data if r["data_type"] == "Academic" and r["metric_2_value"] is not None and float(r["metric_2_value"]) < 50]
    for r in acad_risks:
        ws5.append([
            r["student_id"],
            r["usn"],
            r["student_name"],
            r["course_code"],
            r["course_name"],
            r["batch_name"],
            r["section_name"],
            f"{r['metric_2_value']}%"
        ])
    _format_header(ws5)
    ws5.freeze_panes = "A2"
    _auto_width(ws5)

    # -----------------------------------------
    # 06 Feedback Analysis
    # -----------------------------------------
    ws6 = workbook.create_sheet("06 Feedback Analysis")
    ws6.append(["Course Code", "Course Name", "Module", "Instructor", "Responses", "Average Rating", "Concern (<3/5)"])
    for fb in summary.get("feedback", []):
        is_concern = fb['average_feedback_rating'] is not None and float(fb['average_feedback_rating']) < 3
        ws6.append([
            fb["course_code"],
            fb["course_name"],
            fb["module_name"],
            fb["instructor_name"],
            fb["response_count"],
            f"{fb['average_feedback_rating']} / 5" if fb['average_feedback_rating'] is not None else "N/A",
            "Yes" if is_concern else "No"
        ])
    _format_header(ws6)
    ws6.freeze_panes = "A2"
    _auto_width(ws6)

    # -----------------------------------------
    # 07 HOD Insights
    # -----------------------------------------
    ws7 = workbook.create_sheet("07 HOD Insights")
    ws7.append(["Type", "Insight / Recommendation"])
    for insight in insights.get("insights", []):
        ws7.append(["Insight", insight])
    for recommendation in insights.get("recommendations", []):
        ws7.append(["Recommendation", recommendation])
    _format_header(ws7)
    ws7.column_dimensions["A"].width = 20
    ws7.column_dimensions["B"].width = 100
    for row in ws7.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # -----------------------------------------
    # 08 Data Quality
    # -----------------------------------------
    ws8 = workbook.create_sheet("08 Data Quality")
    att_count = sum(1 for r in raw_data if r["data_type"] == "Attendance")
    acad_count = sum(1 for r in raw_data if r["data_type"] == "Academic")
    feed_count = sum(1 for r in raw_data if r["data_type"] == "Feedback")
    
    status_data = [
        ("Report Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Reporting Period", f"{week_start} to {week_end}"),
        ("Department", department["department_name"]),
        ("Attendance Records", att_count),
        ("Academic Records", acad_count),
        ("Feedback Records", feed_count),
        ("Students Processed", dept["student_count"]),
        ("Courses Processed", dept["course_count"]),
        ("Batches Processed", len(summary.get("batches", []))),
        ("Sections Processed", len(summary.get("sections", []))),
        ("Facilitators Processed", len(summary.get("facilitators", []))),
        ("Academic Data Available", "Yes" if acad_count > 0 else "No"),
        ("Feedback Data Available", "Yes" if feed_count > 0 else "No"),
    ]
    ws8.append(["Metric", "Value"])
    for k, v in status_data:
        ws8.append([k, v])
    _format_header(ws8)
    _auto_width(ws8)

    # -----------------------------------------
    # SAVE
    # -----------------------------------------
    workbook.save(filepath)
    return str(filepath)
